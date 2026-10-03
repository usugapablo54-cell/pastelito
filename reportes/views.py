from django.shortcuts import render
from django.db.models import Sum
from gastos.models import Gasto
from materia_prima.models import Insumo
from ventas.models import Producto
from inventario.models import MovimientoInventario, TipoMovimiento

def reportes(request):
    gastos_por_categoria = (
        Gasto.objects.values("categoria")
        .annotate(total=Sum("monto"))
        .order_by("-total")
    )

    compras_por_proveedor = (
        Insumo.objects.exclude(proveedor__isnull=True)
        .values("proveedor__nombre")
        .annotate(total=Sum("precio_compra"))
        .order_by("-total")
    )

    consumo_materia_prima = (
        MovimientoInventario.objects.filter(tipo=TipoMovimiento.SALIDA)
        .values("insumo__nombre")
        .annotate(total=Sum("cantidad"))
        .order_by("-total")
    )

    productos = Producto.objects.all()
    productos_rentables = sorted(productos, key=lambda p: p.utilidad_unitaria, reverse=True)

    total_gastos = Gasto.objects.aggregate(total=Sum("monto"))["total"] or 0
    total_ventas_estimadas = sum(p.precio_sugerido for p in productos)
    utilidad_estimada = total_ventas_estimadas - total_gastos

    context = {
        "gastos_por_categoria": gastos_por_categoria,
        "compras_por_proveedor": compras_por_proveedor,
        "consumo_materia_prima": consumo_materia_prima,
        "productos_rentables": productos_rentables,
        "total_gastos": total_gastos,
        "total_ventas_estimadas": total_ventas_estimadas,
        "utilidad_estimada": utilidad_estimada,
    }
    return render(request, "reportes/reportes.html", context)
from django.http import HttpResponse
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import openpyxl


def exportar_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="reporte.pdf"'

    p = canvas.Canvas(response, pagesize=letter)
    p.setFont("Helvetica-Bold", 16)
    p.drawString(50, 750, "Reporte General del Negocio")

    p.setFont("Helvetica", 12)
    y = 710

    total_gastos = Gasto.objects.aggregate(total=Sum('monto'))['total'] or 0
    p.drawString(50, y, f"Total gastos: $ {total_gastos:.0f}")
    y -= 30

    p.setFont("Helvetica-Bold", 12)
    p.drawString(50, y, "Gastos por categoria:")
    y -= 20
    p.setFont("Helvetica", 11)
    for g in Gasto.objects.values('categoria').annotate(total=Sum('monto')).order_by('-total'):
        p.drawString(60, y, f"{g['categoria']}: $ {g['total']:.0f}")
        y -= 18

    y -= 15
    p.setFont("Helvetica-Bold", 12)
    p.drawString(50, y, "Productos mas rentables:")
    y -= 20
    p.setFont("Helvetica", 11)
    productos = sorted(Producto.objects.all(), key=lambda x: x.utilidad_unitaria, reverse=True)
    for prod in productos:
        p.drawString(60, y, f"{prod.nombre}: $ {prod.utilidad_unitaria:.2f} c/u")
        y -= 18

    p.showPage()
    p.save()
    return response


def exportar_excel(request):
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Gastos"
    ws1.append(["Categoria", "Descripcion", "Monto", "Fecha"])
    for g in Gasto.objects.all().order_by('-fecha'):
        ws1.append([g.get_categoria_display(), g.descripcion, float(g.monto), str(g.fecha)])

    ws2 = wb.create_sheet("Materia Prima")
    ws2.append(["Nombre", "Precio compra", "Cantidad", "Unidad", "Stock actual", "Valor stock"])
    for i in Insumo.objects.all():
        ws2.append([i.nombre, float(i.precio_compra), float(i.cantidad), i.unidad_medida, float(i.stock_actual), float(i.valor_stock)])

    ws3 = wb.create_sheet("Productos")
    ws3.append(["Producto", "Costo total", "Margen %", "Precio sugerido", "Utilidad"])
    for prod in Producto.objects.all():
        ws3.append([prod.nombre, float(prod.costo_total_unidad), float(prod.margen_ganancia), float(prod.precio_sugerido), float(prod.utilidad_unitaria)])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="reporte.xlsx"'
    wb.save(response)
    return response